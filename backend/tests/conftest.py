import os

# Must happen before `app.main` (and therefore `app.config`) is imported —
# the rate limiter reads RATE_LIMIT once at import time, and the test suite
# fires far more than 30 requests/minute against one shared TestClient "IP".
os.environ.setdefault("RATE_LIMIT", "10000/minute")

import fitz  # PyMuPDF
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation

from app.main import app


@pytest.fixture(scope="session")
def client():
    return TestClient(app)


@pytest.fixture(scope="session")
def fixtures_dir(tmp_path_factory):
    return tmp_path_factory.mktemp("riqopdf-fixtures")


@pytest.fixture(scope="session")
def sample_pdf(fixtures_dir):
    """A plain 3-page text PDF — the default input for most tools."""
    path = fixtures_dir / "sample.pdf"
    doc = fitz.open()
    for i in range(3):
        doc.new_page().insert_text((72, 72), f"Test page {i + 1}\nHello RiqoPDF testing.")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture(scope="session")
def second_pdf(fixtures_dir):
    """A second, distinct PDF for merge."""
    path = fixtures_dir / "second.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Second document page 1")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture(scope="session")
def encrypted_pdf(fixtures_dir):
    path = fixtures_dir / "encrypted.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "secret content")
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="correcthorse", user_pw="correcthorse")
    doc.close()
    return path


@pytest.fixture(scope="session")
def scanned_pdf(fixtures_dir):
    """An image-only PDF (no text layer) for exercising OCR."""
    img_path = fixtures_dir / "scan_page.jpg"
    img = Image.new("RGB", (850, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((100, 100), "SCANNED DOCUMENT", fill="black")
    d.text((100, 150), "This is a test scan for OCR verification.", fill="black")
    img.save(img_path)

    path = fixtures_dir / "scan.pdf"
    doc = fitz.open()
    page = doc.new_page(width=850, height=1100)
    page.insert_image(page.rect, filename=str(img_path))
    doc.save(path)
    doc.close()
    return path


@pytest.fixture(scope="session")
def sample_jpg(fixtures_dir):
    path = fixtures_dir / "sample.jpg"
    img = Image.new("RGB", (400, 300), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([20, 20, 380, 280], outline="black", width=3)
    d.text((150, 140), "Test JPG", fill="black")
    img.save(path)
    return path


@pytest.fixture(scope="session")
def sample_docx(fixtures_dir):
    path = fixtures_dir / "sample.docx"
    doc = Document()
    doc.add_heading("Test Doc")
    doc.add_paragraph("Hello from RiqoPDF testing.")
    doc.save(path)
    return path


@pytest.fixture(scope="session")
def sample_xlsx(fixtures_dir):
    path = fixtures_dir / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Hello"
    ws["B1"] = "RiqoPDF"
    wb.save(path)
    return path


@pytest.fixture(scope="session")
def sample_pptx(fixtures_dir):
    path = fixtures_dir / "sample.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Test Slide"
    prs.save(path)
    return path


@pytest.fixture(scope="session")
def image_heavy_pdf(fixtures_dir):
    """A PDF with real image content, for compress/target-size tests where
    a text-only PDF has nothing to compress."""
    import random

    img_path = fixtures_dir / "noisy.jpg"
    img = Image.new("RGB", (800, 600))
    pixels = img.load()
    for x in range(800):
        for y in range(600):
            pixels[x, y] = (random.randint(0, 255), random.randint(0, 255), random.randint(0, 255))
    img.save(img_path, quality=95)

    path = fixtures_dir / "image_heavy.pdf"
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page(width=800, height=600)
        page.insert_image(page.rect, filename=str(img_path))
    doc.save(path)
    doc.close()
    return path
