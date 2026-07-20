import fitz  # PyMuPDF
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from starlette.background import BackgroundTask

from ..storage import delete_job_dir, new_job_dir
from ..utils import require_pdf, save_upload

router = APIRouter()


@router.post("/pdf-to-excel")
async def pdf_to_excel(file: UploadFile = File(...)):
    """Detects tables per page where possible; otherwise falls back to one text line per row."""
    require_pdf(file.filename)
    job_dir = new_job_dir()
    src = job_dir / "in.pdf"
    await save_upload(file, src)

    try:
        doc = fitz.open(str(src))
        wb = Workbook()
        wb.remove(wb.active)

        for i, page in enumerate(doc):
            sheet = wb.create_sheet(title=f"Page {i + 1}"[:31])
            tables = page.find_tables()
            if tables.tables:
                row_offset = 0
                for table in tables.tables:
                    for row in table.extract():
                        row_offset += 1
                        for col, value in enumerate(row, start=1):
                            sheet.cell(row=row_offset, column=col, value=value)
                    row_offset += 1  # blank row between tables
            else:
                for row_idx, line in enumerate(page.get_text().splitlines(), start=1):
                    if line.strip():
                        sheet.cell(row=row_idx, column=1, value=line)
        doc.close()

        if not wb.sheetnames:
            wb.create_sheet(title="Page 1")

        out_path = job_dir / "converted.xlsx"
        wb.save(out_path)
    except Exception as exc:
        delete_job_dir(job_dir)
        raise HTTPException(400, f"Could not convert PDF to Excel: {exc}") from exc

    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="converted.xlsx",
        background=BackgroundTask(delete_job_dir, job_dir),
    )
