import fitz  # PyMuPDF
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt
from starlette.background import BackgroundTask

from ..storage import delete_job_dir, new_job_dir
from ..utils import require_pdf, save_upload

router = APIRouter()

EMU_PER_POINT = 12700  # 1 pt = 1/72 inch = 12700 EMU


def _add_text_line(slide, line: dict):
    text = "".join(span["text"] for span in line["spans"]).strip()
    if not text:
        return
    x0, y0, x1, y1 = line["bbox"]
    width = max(x1 - x0, 1)
    height = max(y1 - y0, 1)
    box = slide.shapes.add_textbox(
        Emu(int(x0 * EMU_PER_POINT)),
        Emu(int(y0 * EMU_PER_POINT)),
        Emu(int(width * EMU_PER_POINT)),
        Emu(int(height * EMU_PER_POINT)),
    )
    frame = box.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0

    span = line["spans"][0]
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(max(round(span["size"]), 1))
    color = span.get("color", 0)
    run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)


def _add_page_images(doc: fitz.Document, page: fitz.Page, slide, job_dir):
    for i, img in enumerate(page.get_images(full=True)):
        xref = img[0]
        rects = page.get_image_rects(xref)
        if not rects:
            continue
        try:
            base = doc.extract_image(xref)
        except Exception:
            continue
        img_path = job_dir / f"img_{page.number}_{i}.{base['ext']}"
        img_path.write_bytes(base["image"])
        for rect in rects:
            slide.shapes.add_picture(
                str(img_path),
                Emu(int(rect.x0 * EMU_PER_POINT)),
                Emu(int(rect.y0 * EMU_PER_POINT)),
                width=Emu(int(rect.width * EMU_PER_POINT)),
                height=Emu(int(rect.height * EMU_PER_POINT)),
            )


@router.post("/pdf-to-powerpoint")
async def pdf_to_powerpoint(file: UploadFile = File(...)):
    """Rebuilds each page as real, editable text boxes (from PyMuPDF's per-line
    text layout) plus the page's embedded images at their original positions —
    no full-page rasterization, so text stays text."""
    require_pdf(file.filename)
    job_dir = new_job_dir()
    src = job_dir / "in.pdf"
    await save_upload(file, src)

    try:
        doc = fitz.open(str(src))
        prs = Presentation()
        blank_layout = prs.slide_layouts[6]

        first_rect = doc[0].rect if doc.page_count else fitz.Rect(0, 0, 612, 792)
        prs.slide_width = Emu(int(first_rect.width * EMU_PER_POINT))
        prs.slide_height = Emu(int(first_rect.height * EMU_PER_POINT))

        for page in doc:
            slide = prs.slides.add_slide(blank_layout)
            _add_page_images(doc, page, slide, job_dir)
            for block in page.get_text("dict")["blocks"]:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    _add_text_line(slide, line)
        doc.close()

        out_path = job_dir / "converted.pptx"
        prs.save(out_path)
    except Exception as exc:
        delete_job_dir(job_dir)
        raise HTTPException(400, f"Could not convert PDF to PowerPoint: {exc}") from exc

    return FileResponse(
        out_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename="converted.pptx",
        background=BackgroundTask(delete_job_dir, job_dir),
    )
